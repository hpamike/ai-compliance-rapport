#!/usr/bin/env python3
"""Zet een rapport in eenvoudige Markdown om naar een Word-bestand (.docx) in AI FiT-huisstijl.

Gebruik:
    python build_docx.py rapport.md rapport.docx                   # Nederlands, AI FiT-huisstijl
    python build_docx.py rapport.md rapport.docx --en              # Engelse voettekst
    python build_docx.py rapport.md rapport.docx --footer "tekst"  # eigen voettekst (bijv. andere taal)
    python build_docx.py rapport.md rapport.docx --neutraal        # zonder logo en AI FiT-oproep in de koptekst

Ondersteunde Markdown:
    # Titel                 -> documenttitel (paars, vet)
    ## Kop                  -> hoofdstuk (Kop 1, paars, nieuwe pagina vanaf het tweede hoofdstuk)
    ### Subkop              -> tussenkop (Kop 2, paars)
    #### Subsubkop          -> Kop 3 (zwart, vet)
    gewone tekst            -> alinea (lege regel = nieuwe alinea)
    - punt / * punt         -> opsomming (2 of 4 spaties inspringen = dieper niveau)
    1. punt                 -> genummerde lijst (idem)
    | a | b |  (met |---|)  -> tabel; eerste rij is de kopregel
    > tekst                 -> paars kader; "> **LABEL:** tekst" zet LABEL (hoofdletters) als kaderkop
    **vet** en *cursief*    -> binnen een regel
    **Advies:** ...         -> regel die met "**Advies:**" of "**Recommendation:**" begint, wordt paars

Huisstijl (AI FiT): Calibri 10,5 pt, tekst #1A1A1A, regelafstand 1,15; koppen vet paars #6E42D8;
kaders met paarse rand en vulling #EFEAFB; A4 met smalle marges.
Koptekst: logo, oproep tot de gratis AI FiT scan op https://aifit.nu en een teal streep. Voettekst: disclaimer over mogelijke fouten door AI, en "Pagina X van Y".
Vereist: python-docx (pip install python-docx).
"""
import os
import re
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches, Twips
except ImportError:
    sys.exit("python-docx ontbreekt: installeer met 'pip install python-docx' en probeer opnieuw.")

FONT = "Calibri"
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
PURPLE = RGBColor(0x6E, 0x42, 0xD8)
GREY = RGBColor(0x80, 0x80, 0x80)
PURPLE_HEX, FILL_HEX, LINE_HEX, TEAL_HEX = "6E42D8", "EFEAFB", "C9BCF2", "82C3C0"
CTA = {"nl": "Doe de gratis AI FiT scan op https://aifit.nu", "en": "Take the free AI FiT scan at https://aifit.nu"}
LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "aifit-logo.png")
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
PURPLE_LEADS = ("**Advies:**", "**Recommendation:**", "**Conclusie:**", "**Conclusion:**")
FOOTER = {
    "nl": ("Opgesteld met behulp van AI: fouten zijn mogelijk. Geen juridisch advies; "
           "laat altijd controleren door een specialist.", "Pagina ", " van "),
    "en": ("Prepared with the help of AI: errors are possible. Not legal advice; "
           "always have it reviewed by a specialist.", "Page ", " of "),
}


def set_font(style, size, bold=False, color=TEXT):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(theme), None)


def setup(doc, neutral, lang, footer_text):
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.top_margin = sec.bottom_margin = Twips(900)
    sec.left_margin = sec.right_margin = Twips(1200)
    text_width = sec.page_width - sec.left_margin - sec.right_margin

    # documenttaal (spellingscontrole) in de docDefaults, zodat ook kop- en voettekst meedoen
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    for old in rpr_default.findall(qn("w:lang")):
        rpr_default.remove(old)
    lang_el = OxmlElement("w:lang")
    lang_el.set(qn("w:val"), "en-GB" if lang == "en" else "nl-NL")
    rpr_default.append(lang_el)

    normal = doc.styles["Normal"]
    set_font(normal, 10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    set_font(doc.styles["Title"], 16, bold=True, color=PURPLE)
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    for bdr in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(bdr)
    doc.styles["Title"].paragraph_format.space_before = Pt(6)
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    for name, size, color, before, after in (("Heading 1", 13.5, PURPLE, 3, 9),
                                              ("Heading 2", 10.5, PURPLE, 8, 5),
                                              ("Heading 3", 10.5, TEXT, 6, 3)):
        st = doc.styles[name]
        set_font(st, size, bold=True, color=color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Number", "List Number 2", "List Number 3"):
        doc.styles[name].paragraph_format.space_after = Pt(3)

    # Koptekst: logo links, oproep tot de AI FiT scan rechts, teal streep eronder
    header = sec.header.paragraphs[0]
    clear_default_tabs(header)
    header.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    if not neutral:
        header.paragraph_format.space_after = Pt(3)
        if os.path.exists(LOGO):
            header.add_run().add_picture(LOGO, width=Inches(0.34))
        r = header.add_run("\t" + CTA[lang])
        r.font.size = Pt(9)
        r.font.color.rgb = PURPLE
        line = sec.header.add_paragraph()
        line.paragraph_format.space_after = Pt(6)
        ppr = line._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "1"), ("w:color", TEAL_HEX)):
            bottom.set(qn(k), v)
        pbdr.append(bottom)
        insert_before(ppr, pbdr, ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc", "w:rPr"))

    # Voettekst: disclaimer links, paginanummer rechts
    disclaimer, page, of = FOOTER[lang]
    footer = sec.footer.paragraphs[0]
    clear_default_tabs(footer)
    footer.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    grey_run(footer, (footer_text or disclaimer) + "\t" + page)
    add_field(footer, "PAGE")
    grey_run(footer, of)
    add_field(footer, "NUMPAGES")


def clear_default_tabs(par):
    """De standaardstijlen Header en Footer hebben tabs op 3,25 en 6,5 inch; die zitten de rechter tab in de weg."""
    for pos in (Twips(4680), Twips(9360)):
        par.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)


def grey_run(par, text):
    r = par.add_run(text)
    r.font.size = Pt(7.5)
    r.font.color.rgb = GREY
    return r


def add_field(par, code):
    run = par.add_run()
    run.font.size = Pt(7.5)
    run.font.color.rgb = GREY
    for kind, text in (("begin", None), (None, code), ("separate", None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def add_inline(par, text, color=None, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.italic = True
        else:
            run = par.add_run(part)
        if color:
            run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)


def insert_before(parent, child, followers):
    """Voeg child in vóór het eerste element uit followers (het OOXML-schema eist een vaste volgorde)."""
    for el in parent:
        if el.tag in [qn(f) for f in followers]:
            el.addprevious(child)
            return
    parent.append(child)


def borders(el_pr, tag, color, size):
    b = OxmlElement(tag)
    for edge in ("top", "left", "bottom", "right") + (("insideH", "insideV") if tag == "w:tblBorders" else ()):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    if tag == "w:tblBorders":
        insert_before(el_pr, b, ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"))
    else:
        insert_before(el_pr, b, ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"))


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table._tbl.tblPr, "w:tblBorders", LINE_HEX, 4)
    for i, row in enumerate(cells):
        for j in range(width):
            cell = table.cell(i, j)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing = 1.0
            add_inline(par, row[j] if j < len(row) else "", size=9.5)
            if i == 0:
                shade(cell, FILL_HEX)
                for run in par.runs:
                    run.bold = True
                    run.font.color.rgb = PURPLE
    doc.add_paragraph()


def add_box(doc, text):
    """Paars kader zoals het promptkader: rand #6E42D8, vulling #EFEAFB, optioneel label in hoofdletters."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    borders(cell._tc.get_or_add_tcPr(), "w:tcBorders", PURPLE_HEX, 8)
    shade(cell, FILL_HEX)
    par = cell.paragraphs[0]
    m = re.match(r"\*\*([^*]+?):?\*\*:?\s*(.*)", text)
    if m and m.group(1).isupper():
        label = par.add_run(m.group(1))
        label.bold = True
        label.font.size = Pt(10)
        label.font.color.rgb = PURPLE
        par = cell.add_paragraph()
        text = m.group(2)
    add_inline(par, text, size=10)
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
    doc.add_paragraph()


def build(md_text, out_path, neutral=False, lang="nl", footer_text=None):
    doc = Document()
    setup(doc, neutral, lang, footer_text)
    lines = md_text.splitlines()
    para, i, chapters = [], 0, 0

    def flush():
        if para:
            text = " ".join(para)
            p = doc.add_paragraph()
            add_inline(p, text, color=PURPLE if text.startswith(PURPLE_LEADS) else None)
            para.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            flush()
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        elif s.startswith(">"):
            flush()
            box = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                box.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_box(doc, " ".join(b for b in box if b))
            continue
        elif s.startswith("#### "):
            flush(); doc.add_heading(s[5:], level=3)
        elif s.startswith("### "):
            flush(); doc.add_heading(s[4:], level=2)
        elif s.startswith("## "):
            flush()
            h = doc.add_heading(s[3:], level=1)
            chapters += 1
            if chapters > 1:
                h.paragraph_format.page_break_before = True
        elif s.startswith("# "):
            flush(); doc.add_heading(s[2:], level=0)
        elif re.match(r"^[-*] ", s):
            flush()
            level = min((len(line) - len(line.lstrip())) // 2, 2)
            style = "List Bullet" if level == 0 else f"List Bullet {level + 1}"
            add_inline(doc.add_paragraph(style=style), s[2:])
        elif re.match(r"^\d+[.)] ", s):
            flush()
            level = min((len(line) - len(line.lstrip())) // 2, 2)
            style = "List Number" if level == 0 else f"List Number {level + 1}"
            add_inline(doc.add_paragraph(style=style), re.sub(r"^\d+[.)] ", "", s))
        else:
            para.append(s)
        i += 1
    flush()
    doc.save(out_path)


if __name__ == "__main__":
    argv = sys.argv[1:]
    footer_text = None
    if "--footer" in argv:
        k = argv.index("--footer")
        footer_text = argv[k + 1]
        del argv[k:k + 2]
    args = [a for a in argv if not a.startswith("--")]
    if len(args) != 2:
        sys.exit(__doc__)
    with open(args[0], encoding="utf-8") as f:
        build(f.read(), args[1],
              neutral="--neutraal" in argv or "--neutral" in argv,
              lang="en" if "--en" in argv else "nl",
              footer_text=footer_text)
    print(f"geschreven: {args[1]}")
